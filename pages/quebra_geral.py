"""
quebra.py
=========
Super Relatório Corporativo Unificado | Quebra Operacional TOTALE

Versão: 5.4.0 (Busca Avançada no Google Sheets + Conversão Login -> Nome do Técnico/Monitor)
Author: TOTALE Tecnologia
"""

from __future__ import annotations

import csv
import hashlib
import importlib
import logging
import re
import sys
import unicodedata
import urllib.request
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import datetime
from html import escape
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Final, Literal, cast
from urllib.parse import quote

import numpy as np
import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ═════════════════════════════════════════════════════════════════════
# PATH BOOTSTRAP (antes de qualquer import local)
# ═════════════════════════════════════════════════════════════════════
_DIR: Final[Path] = Path(__file__).resolve().parent
_ROOT: Final[Path] = _DIR.parent
for _path in (_DIR, _ROOT):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))

logger = logging.getLogger(__name__)

# ═════════════════════════════════════════════════════════════════════
# TIPOS
# ═════════════════════════════════════════════════════════════════════
TemaKPI = Literal[
    "azul", "verde", "vermelho", "laranja", "cinza", "roxo", "amarelo", "escuro"
]
TipoInsight = Literal["ok", "info", "alerta", "critico", "acao"]


# ═════════════════════════════════════════════════════════════════════
# IMPORTS OPCIONAIS – CRITÉRIOS
# ═════════════════════════════════════════════════════════════════════
@dataclass
class ModuloCriterios:
    """Encapsula o módulo opcional de critérios."""

    disponivel: bool = False
    termos_nd: tuple[str, ...] = ("ADESAO", "ADESÃO")
    termo_gpon: str = "PON"
    termo_migracao: str = "MUDANCA DE PACOTE"
    termo_pme: str = "PME"
    vazios_gerais: set[str] = field(default_factory=set)
    classificar: Callable[..., Any] | None = None
    criar_flag_gpon: Callable[..., Any] | None = None
    detect_tipo_os: Callable[..., Any] | None = None
    detect_habilidade: Callable[..., Any] | None = None
    detect_flag_gpon: Callable[..., Any] | None = None
    render_painel: Callable[..., Any] | None = None


def _carregar_criterios() -> ModuloCriterios:
    try:
        from components import criterios as c  # type: ignore

        return ModuloCriterios(
            disponivel=True,
            termos_nd=c.TERMOS_ND,
            termo_gpon=c.TERMO_GPON_HABILIDADE,
            termo_migracao=c.TERMO_MIGRACAO_OS,
            termo_pme=c.TERMO_PME_HABILIDADE,
            vazios_gerais=c.VAZIOS_GERAIS,
            classificar=c.classificar_tipo_servico,
            criar_flag_gpon=c.criar_flag_gpon,
            detect_tipo_os=c.detectar_col_tipo_os_1,
            detect_habilidade=c.detectar_col_habilidade,
            detect_flag_gpon=c.detectar_col_flag_gpon,
            render_painel=c.render_painel_criterios,
        )
    except ImportError as erro:
        logger.warning("Módulo critérios não disponível: %s", erro)
        return ModuloCriterios()


CRITERIOS: Final[ModuloCriterios] = _carregar_criterios()


# ═════════════════════════════════════════════════════════════════════
# IMPORTS OPCIONAIS – COMPONENTES VISUAIS
# ═════════════════════════════════════════════════════════════════════
@dataclass
class ComponentesVisuais:
    """Componentes de UI opcionais."""

    disponivel: bool = False
    section_header: Callable[..., Any] | None = None
    sidebar_brand: Callable[..., Any] | None = None
    table_html: Callable[..., Any] | None = None
    insight: Callable[..., Any] | None = None
    kpi: Callable[..., Any] | None = None
    kpi_sm: Callable[..., Any] | None = None
    theme_selector: Callable[..., Any] | None = None


def _carregar_componentes() -> ComponentesVisuais:
    try:
        from components import componentes as m  # type: ignore

        return ComponentesVisuais(
            disponivel=True,
            section_header=getattr(m, "render_section_header", None),
            sidebar_brand=getattr(m, "render_sidebar_brand", None),
            table_html=getattr(m, "render_table_html", None),
            insight=getattr(m, "render_insight", None),
            kpi=getattr(m, "render_kpi", None),
            kpi_sm=getattr(m, "render_kpi_sm", None),
            theme_selector=getattr(m, "render_page_sidebar_theme_selector", None),
        )
    except ImportError as erro:
        logger.warning("Módulo componentes não disponível: %s", erro)
        return ComponentesVisuais()


COMPONENTES: Final[ComponentesVisuais] = _carregar_componentes()


# ═════════════════════════════════════════════════════════════════════
# IMPORTS OPCIONAIS – ROBÔ
# ═════════════════════════════════════════════════════════════════════
def _carregar_robo() -> tuple[Callable[..., Any] | None, str]:
    try:
        module = importlib.import_module("robo.robo_local")
        funcao = getattr(module, "renderizar_robo_local", None) or getattr(
            module, "render_robo_local", None
        )
        if callable(funcao):
            return funcao, f"OK ({getattr(module, '__file__', '?')})"
        return None, "Função de renderização não encontrada."
    except Exception as erro:
        logger.debug("Robô local não pôde ser importado.", exc_info=True)
        return None, f"{type(erro).__name__}: {erro}"


_ROBO_FN, _ROBO_ERRO = _carregar_robo()
ROBO_DISPONIVEL: bool = _ROBO_FN is not None


# ═════════════════════════════════════════════════════════════════════
# CONFIGURAÇÕES
# ═════════════════════════════════════════════════════════════════════
class Config:
    """Constantes de negócio e visuais."""

    SLA_QUEBRA_MAXIMA: Final[float] = 0.20
    SLA_MIGRACAO_MAXIMA: Final[float] = 0.25

    URL_LISTA_ATIVOS: Final[str] = (
        "https://docs.google.com/spreadsheets/d/"
        "1LQKDcLshC6XSXLBVWaEYSpxrro6uydyU9pwDLc38pEg/edit"
    )
    SHEET_ID_ATIVOS: Final[str] = "1LQKDcLshC6XSXLBVWaEYSpxrro6uydyU9pwDLc38pEg"
    WORKSHEET_ATIVOS: Final[str] = "lista_ativos"

    STATUS_ORDEM: Final[tuple[str, ...]] = ("Executada", "Não Executada", "Pendente")
    ORDEM_TIPOS: Final[tuple[str, ...]] = (
        "Novos Domicílios",
        "PME",
        "Migração",
        "Outros",
    )

    COLUNAS_TIPO_SERVICO: Final[tuple[str, ...]] = (
        "TIPO_SERVICO",
        "TIPO SERVICO",
        "TIPO O.S 1",
        "TIPO OS 1",
        "TIPO O.S. 1",
    )

    CORES_STATUS: Final[dict[str, str]] = {
        "Executada": "#10B981",
        "Não Executada": "#EF4444",
        "Pendente": "#94A3B8",
    }
    CORES_TIPO: Final[dict[str, str]] = {
        "Novos Domicílios": "#1E40AF",
        "PME": "#7C3AED",
        "Migração": "#0369A1",
        "Quebra Geral": "#78350F",
        "Outros": "#64748B",
    }


class Fontes:
    TITULO: Final[str] = "'Manrope', sans-serif"
    TEXTO: Final[str] = "'Inter', sans-serif"


# ═════════════════════════════════════════════════════════════════════
# MAPEAMENTOS
# ═════════════════════════════════════════════════════════════════════
_CODIGOS_NAO_EXECUTADA: Final[tuple[int, ...]] = (
    100,
    101,
    103,
    104,
    105,
    106,
    107,
    108,
    110,
    112,
    113,
    114,
    125,
    203,
    204,
    205,
    206,
    301,
    302,
    303,
    305,
    306,
    307,
    308,
    312,
    316,
    400,
    402,
)
_CODIGOS_EXECUTADA: Final[tuple[int, ...]] = (
    128,
    328,
    408,
    409,
    425,
    430,
    440,
    467,
    470,
    474,
    475,
    477,
    *range(500, 591),
)
MAPA_CODIGO_NUMERICO: Final[dict[int, str]] = {
    **{c: "Não Executada" for c in _CODIGOS_NAO_EXECUTADA},
    **{c: "Executada" for c in _CODIGOS_EXECUTADA},
}

CORES_REGIAO: Final[dict[str, dict[str, str]]] = {
    "LESTE": {"bg": "#DBEAFE", "text": "#1E40AF", "border": "#3B82F6"},
    "GRU": {"bg": "#D1FAE5", "text": "#065F46", "border": "#10B981"},
    "ABCDM": {"bg": "#EDE9FE", "text": "#5B21B6", "border": "#8B5CF6"},
    "OUTRAS": {"bg": "#F1F5F9", "text": "#475569", "border": "#94A3B8"},
}

CIDADES_LESTE: Final[frozenset[str]] = frozenset({"SAO PAULO"})
CIDADES_GRU: Final[frozenset[str]] = frozenset(
    {
        "GUARULHOS",
        "ARUJA",
        "MOGI DAS CRUZES",
        "SUZANO",
        "ITAQUAQUECETUBA",
        "FERRAZ DE VASCONCELOS",
        "POA",
    }
)
CIDADES_ABCDM: Final[frozenset[str]] = frozenset(
    {
        "SANTO ANDRE",
        "SAO BERNARDO DO CAMPO",
        "SAO CAETANO DO SUL",
        "DIADEMA",
        "MAUA",
        "RIBEIRAO PIRES",
        "RIO GRANDE DA SERRA",
    }
)

VALORES_AUSENTES: Final[frozenset[str]] = frozenset(
    {
        "",
        "NAN",
        "NONE",
        "NULL",
        "NAT",
        "NA",
        "N/A",
        "-",
        "None",
    }
)

RE_ESPACOS: Final[re.Pattern[str]] = re.compile(r"\s+")
RE_ALFANUMERICO: Final[re.Pattern[str]] = re.compile(r"[^A-Z0-9]")
RE_NUMERO_INICIAL: Final[re.Pattern[str]] = re.compile(r"^(\d+)")
RE_LOGIN_DECIMAL: Final[re.Pattern[str]] = re.compile(r"^(\d+)\.0+$")


# ═════════════════════════════════════════════════════════════════════
# HELPERS DE FORMATAÇÃO
# ═════════════════════════════════════════════════════════════════════
def html_safe(valor: Any) -> str:
    return escape(str(valor), quote=True)


def fmt_pct_br(valor: Any) -> str:
    try:
        n = float(valor)
        if not np.isfinite(n):
            return "—"
        return f"{n * 100:,.2f}%".replace(",", "X").replace(".", ",").replace("X", ".")
    except (TypeError, ValueError):
        return "—"


def fmt_int_br(valor: Any) -> str:
    try:
        n = float(valor)
        if not np.isfinite(n):
            return "—"
        return f"{int(round(n)):,}".replace(",", ".")
    except (TypeError, ValueError):
        return "—"


def is_missing(valor: Any) -> bool:
    try:
        r = pd.isna(valor)
        return bool(r) if isinstance(r, (bool, np.bool_)) else False
    except (TypeError, ValueError):
        return False


def div_segura(num: pd.Series, den: pd.Series, padrao: float = 0.0) -> pd.Series:
    n = pd.to_numeric(num, errors="coerce").fillna(0).to_numpy(dtype=float)
    d = pd.to_numeric(den, errors="coerce").fillna(0).to_numpy(dtype=float)
    out = np.full_like(n, padrao, dtype=float)
    np.divide(n, d, out=out, where=d > 0)
    return pd.Series(out, index=num.index)


def hash_upload(nome: str, conteudo: bytes) -> str:
    return f"{nome}|{len(conteudo)}|{hashlib.sha256(conteudo).hexdigest()}"


def df_sessao(chave: str) -> pd.DataFrame | None:
    v = st.session_state.get(chave)
    return v if isinstance(v, pd.DataFrame) else None


# ═════════════════════════════════════════════════════════════════════
# NORMALIZAÇÃO DE TEXTO E TIPOS
# ═════════════════════════════════════════════════════════════════════
def normalizar_texto(valor: Any) -> str:
    if valor is None:
        return ""
    try:
        if isinstance(pd.isna(valor), (bool, np.bool_)) and bool(pd.isna(valor)):
            return ""
    except (TypeError, ValueError):
        pass
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return RE_ESPACOS.sub(" ", texto.upper().strip())


def normalizar_coluna(valor: Any) -> str:
    return RE_ALFANUMERICO.sub("", normalizar_texto(valor))


def normalizar_login(valor: Any) -> str:
    if valor is None:
        return ""
    texto = normalizar_texto(valor).replace(" ", "")
    m = RE_LOGIN_DECIMAL.match(texto)
    if m:
        texto = m.group(1)
    return texto


def padronizar_tipo_servico(tipo: Any) -> str:
    if tipo is None or (isinstance(tipo, float) and np.isnan(tipo)):
        return "Outros"

    t = normalizar_texto(tipo)
    if not t:
        return "Outros"

    if any(k in t for k in ("PME", "PJ", "JURIDICO", "CORPORATIVO", "EMPRESA")):
        return "PME"
    if any(k in t for k in ("MIGRA", "MUDANCA DE PACOTE", "UPGRADE", "TROCA", "SWAP")):
        return "Migração"
    if (
        any(
            k in t
            for k in (
                "NOVOS DOMICILIOS",
                "NOVO DOMICILIO",
                "DOMICIL",
                "ADESAO",
                "INSTALACAO",
                "INSTAL",
                "RESIDENCIAL",
                "FIBRA NOVA",
            )
        )
        or t == "ND"
    ):
        return "Novos Domicílios"
    return "Outros"


# ═════════════════════════════════════════════════════════════════════
# UTILITÁRIOS DE DATAFRAME
# ═════════════════════════════════════════════════════════════════════
class DF:
    @staticmethod
    def obter_serie(
        df: pd.DataFrame | None,
        coluna: str | None,
        padrao: Any = "",
        *,
        copiar: bool = False,
    ) -> pd.Series:
        if df is None:
            return pd.Series(dtype="object", name=coluna)
        if coluna is None or coluna not in df.columns:
            return pd.Series(padrao, index=df.index, dtype="object", name=coluna)

        obtido = df[coluna]
        if isinstance(obtido, pd.DataFrame):
            df_obtido = cast(pd.DataFrame, obtido)
            if df_obtido.shape[1] > 1:
                transposta = df_obtido.T
                transposta_preenchida = cast(pd.DataFrame, transposta.bfill())
                serie = transposta_preenchida.iloc[0]
            else:
                serie = df_obtido.T.iloc[0]
            obtido = cast(pd.Series, serie)

        serie_final = cast(pd.Series, obtido)
        if not serie_final.index.equals(df.index):
            serie_final = serie_final.reindex(df.index)
        serie_final.name = coluna
        return serie_final.copy() if copiar else serie_final

    @staticmethod
    def padronizar_colunas(df: pd.DataFrame) -> pd.DataFrame:
        resultado = df.copy()
        nomes, usados = [], {}
        for col in resultado.columns:
            base = normalizar_texto(col) or "COLUNA"
            usados[base] = usados.get(base, 0) + 1
            nomes.append(base if usados[base] == 1 else f"{base}__{usados[base]}")
        resultado.columns = nomes
        return resultado

    @staticmethod
    def buscar_coluna(
        df: pd.DataFrame | None, palavras: tuple[str, ...] | list[str]
    ) -> str | None:
        if df is None:
            return None
        mapa = {normalizar_coluna(c): str(c) for c in df.columns}
        for p in palavras:
            alvo = normalizar_coluna(p)
            if alvo and alvo in mapa:
                return mapa[alvo]
        for p in palavras:
            alvo = normalizar_coluna(p)
            if len(alvo) < 3:
                continue
            for norm, orig in mapa.items():
                if alvo == "CONTRATO" and "STATUS" in norm:
                    continue
                if alvo in norm or norm in alvo:
                    return orig
        return None

    @staticmethod
    def remover_duplicatas_coluna(df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        vistos: dict[str, str] = {}
        a_remover: list[str] = []
        for col in df.columns:
            norm = normalizar_texto(col)
            if norm in vistos:
                a_remover.append(col)
            else:
                vistos[norm] = col
        return df.drop(columns=a_remover) if a_remover else df

    @staticmethod
    def corrigir_colunas_vazias(df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        problema = ["DATA AGENDA MDU", "ID SGD", "DATA_AGENDA_MDU", "ID_SGD"]
        for col in problema:
            if col in df.columns and (
                df[col].isna().all() or (df[col].astype(str) == "None").all()
            ):
                for alt in df.columns:
                    if alt == col:
                        continue
                    up = alt.upper()
                    if ("AGENDA" in up and col in up) or ("ID" in up and "SGD" in up):
                        df[col] = df[alt]
                        break
        return df

    @staticmethod
    def limpar_texto(serie: pd.Series, padrao: str) -> pd.Series:
        r = serie.fillna("").astype(str).str.strip().str.upper()
        return r.mask(r.isin(VALORES_AUSENTES), padrao)


# ═════════════════════════════════════════════════════════════════════
# CONVERSÕES NUMÉRICAS
# ═════════════════════════════════════════════════════════════════════
def converter_numero(valor: Any) -> float:
    if valor is None:
        return np.nan
    if isinstance(valor, (int, float, np.number)):
        n = float(valor)
        return n if np.isfinite(n) else np.nan

    texto = re.sub(r"[^\d,.\-]", "", str(valor).strip())
    if not texto or texto in {"-", ".", ","}:
        return np.nan

    try:
        if "," in texto and "." in texto:
            texto = (
                texto.replace(".", "").replace(",", ".")
                if texto.rfind(",") > texto.rfind(".")
                else texto.replace(",", "")
            )
        elif "," in texto:
            texto = (
                texto.replace(",", "")
                if re.fullmatch(r"-?\d{1,3}(?:,\d{3})+", texto)
                else texto.replace(",", ".")
            )
        elif "." in texto and re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+", texto):
            texto = texto.replace(".", "")
        n = float(texto)
        return n if np.isfinite(n) else np.nan
    except (TypeError, ValueError):
        return np.nan


# ═════════════════════════════════════════════════════════════════════
# CLASSIFICAÇÃO DE STATUS
# ═════════════════════════════════════════════════════════════════════
class ClassificadorStatus:
    @staticmethod
    def status_texto(valor: Any) -> str | None:
        texto = normalizar_texto(valor)
        if texto in VALORES_AUSENTES:
            return None
        if "SUSPENS" in texto:
            return "Suspenso"
        if "CANCEL" in texto:
            return "Cancelado"
        if "NAO EXECUT" in texto or "NAO CONCLU" in texto:
            return "Não Executada"
        if "EXECUT" in texto or "CONCLUID" in texto:
            return "Executada"
        if any(
            t in texto for t in ("PEND", "EM ROTA", "INICIADO", "AGENDADO", "ABERTO")
        ):
            return "Pendente"
        return None

    @staticmethod
    def codigo_baixa(valor: Any) -> str | None:
        texto = normalizar_texto(valor)
        if texto in VALORES_AUSENTES:
            return None
        if texto in {"EM ROTA", "INICIADO", "PENDENTE"}:
            return "Pendente"
        m = RE_NUMERO_INICIAL.match(texto)
        return MAPA_CODIGO_NUMERICO.get(int(m.group(1))) if m else None

    @classmethod
    def classificar(cls, df: pd.DataFrame) -> pd.Series:
        col_inicio = DF.buscar_coluna(
            df, ("INÍCIO", "INICIO", "DATA INÍCIO", "DT INICIO", "HORA INICIO")
        )
        col_fechamento = DF.buscar_coluna(
            df,
            (
                "MOTIVO DE FECHAMENTO EXTERNO",
                "FECHAMENTO EXTERNO",
                "MOTIVO DE FECHAMENTO",
            ),
        )
        col_atividade = DF.buscar_coluna(
            df, ("STATUS DA ATIVIDADE", "STATUS ATIVIDADE", "STATUS_ATIVIDADE")
        )
        col_cod_baixa = DF.buscar_coluna(
            df, ("CÓD DE BAIXA 1", "COD DE BAIXA 1", "MOTIVO DE BAIXA", "COD_BAIXA")
        )
        col_status_os = DF.buscar_coluna(
            df,
            (
                "STATUS DA O.S 1",
                "STATUS OS 1",
                "STATUS CONTRATO",
                "STATUS O.S.",
                "STATUS OS",
            ),
        )

        inicio = DF.obter_serie(df, col_inicio).map(normalizar_texto)
        fechamento = DF.obter_serie(df, col_fechamento).map(normalizar_texto)
        atividade = DF.obter_serie(df, col_atividade).map(normalizar_texto)
        cod_baixa = DF.obter_serie(df, col_cod_baixa)
        status_os = DF.obter_serie(df, col_status_os).map(cls.status_texto)

        s_atividade = atividade.map(cls.status_texto)
        s_codigo = cod_baixa.map(cls.codigo_baixa)

        resultado = s_codigo.where(s_codigo.notna(), status_os)
        resultado = resultado.where(resultado.notna(), s_atividade).fillna("Pendente")

        fech_alvo = {"LIBERADO NO SISTEMA NETSMS", "CANCELADO NO SISTEMA NETSMS"}
        inicio_vazio = inicio.isin(VALORES_AUSENTES)
        c_fech = fechamento.isin(fech_alvo)
        c_cancel = atividade.str.contains("CANCEL", regex=False, na=False)
        c_susp = atividade.str.contains("SUSPENS", regex=False, na=False)
        c_naoconc = atividade.str.contains("NAO CONCLU", regex=False, na=False)

        resultado = resultado.mask(c_naoconc, "Não Executada")
        resultado = resultado.mask(c_susp, "Suspenso")
        resultado = resultado.mask(c_cancel, "Cancelado")
        resultado = resultado.mask(~inicio_vazio & c_fech, "Não Executada")
        resultado = resultado.mask(inicio_vazio & c_fech, "Cancelado")

        return resultado.astype("object")


# ═════════════════════════════════════════════════════════════════════
# GERAÇÃO DE EXCEL
# ═════════════════════════════════════════════════════════════════════
_PALAVRAS_PCT_EXCEL: Final[tuple[str, ...]] = (
    "QUEBRA",
    "FECHAMENTO",
    "CONVERSÃO",
    "CONVERSAO",
    "PROJEÇÃO",
    "PROJECAO",
    "%",
    "PME",
    "MIGRAÇÃO",
    "MIGRACAO",
    "DOMICÍLIOS",
    "DOMICILIOS",
    "GPON",
)


def gerar_excel(df: pd.DataFrame, aba: str = "Dados") -> bytes:
    nome_aba = re.sub(r"[\[\]:*?/\\]", "_", aba).strip()[:31] or "Dados"
    buffer = BytesIO()

    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=nome_aba)
        ws = writer.sheets[nome_aba]

        header_fill = PatternFill("solid", fgColor="0F172A")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_border = Border(bottom=Side(style="thin", color="CBD5E1"))
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = header_border

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for i, col in enumerate(df.columns, start=1):
            nome_upper = str(col).upper()
            amostra = df[col].head(2000).fillna("").astype(str).tolist()
            tamanhos = [len(str(col)), *(len(v) for v in amostra)]
            ws.column_dimensions[get_column_letter(i)].width = min(
                max(max(tamanhos, default=12) + 2, 12), 40
            )

            if any(t in nome_upper for t in _PALAVRAS_PCT_EXCEL):
                for linha in range(2, ws.max_row + 1):
                    ws.cell(linha, i).number_format = "0.00%"

    return buffer.getvalue()


# ═════════════════════════════════════════════════════════════════════
# CLASSE DE COMPATIBILIDADE EXTRA (UTILS)
# ═════════════════════════════════════════════════════════════════════
class Utils:
    """
    Abstração unificada de funções de formatação e processamento
    utilizada por múltiplos módulos da plataforma (ex: quebra_unificada.py).
    """

    @staticmethod
    def classificar_status_excel(df: pd.DataFrame) -> pd.Series:
        """Interface estática para classificar o status contratual da base."""
        return ClassificadorStatus.classificar(df)

    @staticmethod
    def gerar_excel(df: pd.DataFrame, aba: str = "Dados") -> bytes:
        """Interface estática para gerar arquivos Excel binários e formatados."""
        return gerar_excel(df, aba)


# ═════════════════════════════════════════════════════════════════════
# LOADER DE ARQUIVOS
# ═════════════════════════════════════════════════════════════════════
class DataLoader:
    _ENCODINGS_CSV: Final[tuple[str, ...]] = ("utf-8-sig", "utf-8", "latin-1", "cp1252")

    @staticmethod
    @st.cache_data(show_spinner=False)
    def ler_arquivo(file_bytes: bytes, filename: str) -> pd.DataFrame:
        if not file_bytes:
            raise ValueError("O arquivo está vazio.")
        ext = Path(filename).suffix.lower()
        if ext == ".csv":
            return DataLoader._ler_csv(file_bytes)
        if ext in {".xlsx", ".xlsm"}:
            return pd.read_excel(BytesIO(file_bytes), engine="openpyxl", dtype=str)
        raise ValueError("Formato inválido. Utilize CSV, XLSX ou XLSM.")

    @staticmethod
    def _ler_csv(file_bytes: bytes) -> pd.DataFrame:
        texto: str | None = None
        for enc in DataLoader._ENCODINGS_CSV:
            try:
                texto = file_bytes.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if texto is None:
            raise ValueError("Encoding do CSV não identificado.")

        amostra = texto[:10000]
        try:
            delim = csv.Sniffer().sniff(amostra, delimiters=";,\t|").delimiter
        except csv.Error:
            delim = max((";", ",", "\t", "|"), key=amostra.count)

        return pd.read_csv(
            StringIO(texto),
            sep=delim,
            dtype=str,
            engine="python",
            on_bad_lines="skip",
            keep_default_na=False,
        )

    @staticmethod
    @st.cache_data(ttl=300, show_spinner="Sincronizando Lista de Ativos...")
    def buscar_gsheets() -> pd.DataFrame:
        """
        Download multi-estratégia do Google Sheets público:
        1. Exportação CSV direta com User-Agent customizado.
        2. Endpoint GViz/TQ.
        3. Streamlit Connection (caso configurada).
        """
        # Método 1: Export CSV direto (Mais confiável para sheets públicos)
        url_export = (
            f"https://docs.google.com/spreadsheets/d/"
            f"{Config.SHEET_ID_ATIVOS}/export?format=csv&sheet={quote(Config.WORKSHEET_ATIVOS)}"
        )
        try:
            req = urllib.request.Request(
                url_export,
                headers={
                    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
                },
            )
            with urllib.request.urlopen(req, timeout=12) as resp:
                data = resp.read()
            if data:
                raw = pd.read_csv(BytesIO(data), dtype=str)
                if not raw.empty:
                    df_proc = DataLoader.processar_lista_ativos(raw)
                    if not df_proc.empty:
                        return df_proc
        except Exception as e:
            logger.debug("Falha Método 1 (Export CSV): %s", e)

        # Método 2: Endpoint GViz
        url_gviz = (
            f"https://docs.google.com/spreadsheets/d/"
            f"{Config.SHEET_ID_ATIVOS}/gviz/tq?tqx=out:csv&sheet={quote(Config.WORKSHEET_ATIVOS)}"
        )
        try:
            req = urllib.request.Request(
                url_gviz,
                headers={
                    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
                },
            )
            with urllib.request.urlopen(req, timeout=12) as resp:
                data = resp.read()
            if data:
                raw = pd.read_csv(BytesIO(data), dtype=str)
                if not raw.empty:
                    df_proc = DataLoader.processar_lista_ativos(raw)
                    if not df_proc.empty:
                        return df_proc
        except Exception as e:
            logger.debug("Falha Método 2 (GViz): %s", e)

        # Método 3: Streamlit GSheets Connection
        try:
            m = importlib.import_module("streamlit_gsheets")
            conn = st.connection("gsheets", type=m.GSheetsConnection)
            raw = conn.read(
                spreadsheet=Config.URL_LISTA_ATIVOS, worksheet=Config.WORKSHEET_ATIVOS
            )
            if isinstance(raw, pd.DataFrame) and not raw.empty:
                return DataLoader.processar_lista_ativos(raw)
        except Exception as e:
            logger.debug("Falha Método 3 (Streamlit Connection): %s", e)

        return pd.DataFrame()

    @staticmethod
    def processar_lista_ativos(raw: pd.DataFrame) -> pd.DataFrame:
        if raw is None or raw.empty:
            return pd.DataFrame()

        base = raw.copy()
        base.columns = base.columns.astype(str).str.strip()

        aliases: dict[str, tuple[str, ...]] = {
            "LOGIN": (
                "LOGIN",
                "LOGIN DO TÉCNICO",
                "LOGIN DO TECNICO",
                "LOGIN_TECNICO",
                "MATRÍCULA",
                "MATRICULA",
                "ID",
                "USUARIO",
                "USUÁRIO",
                "RE",
                "COD_TECNICO",
            ),
            "TÉCNICO": (
                "TÉCNICO",
                "TECNICO",
                "NOME",
                "COLABORADOR",
                "NOME TÉCNICO",
                "NOME TECNICO",
                "NOME COMPLETO",
                "NOME DO TÉCNICO",
                "NOME DO TECNICO",
            ),
            "MONITOR": (
                "MONITOR",
                "GESTOR",
                "SUPERVISOR",
                "COORDENADOR",
                "LIDER",
                "LÍDER",
                "NOME MONITOR",
                "NOME SUPERVISOR",
                "NOME GESTOR",
            ),
            "BASE": ("BASE", "REGIÃO", "REGIAO", "FILIAL", "LOCALIDADE"),
        }

        resultado = pd.DataFrame(index=base.index)
        for destino, nomes in aliases.items():
            col = DF.buscar_coluna(base, nomes)
            if col:
                resultado[destino] = DF.obter_serie(base, col)

        if "LOGIN" not in resultado.columns:
            return pd.DataFrame()

        resultado["LOGIN"] = resultado["LOGIN"].map(normalizar_login)
        resultado = resultado.loc[~resultado["LOGIN"].isin(VALORES_AUSENTES)].copy()

        # Limpar texto dos nomes para garantir apresentação perfeita
        if "TÉCNICO" in resultado.columns:
            resultado["TÉCNICO"] = (
                resultado["TÉCNICO"].fillna("").astype(str).str.strip().str.upper()
            )
        if "MONITOR" in resultado.columns:
            resultado["MONITOR"] = (
                resultado["MONITOR"].fillna("").astype(str).str.strip().str.upper()
            )

        return resultado.drop_duplicates(subset=["LOGIN"], keep="last").reset_index(
            drop=True
        )

    @staticmethod
    def classificar_regiao(valor: Any) -> str:
        cidade = normalizar_texto(valor)
        if cidade in CIDADES_LESTE:
            return "LESTE"
        if cidade in CIDADES_GRU:
            return "GRU"
        if cidade in CIDADES_ABCDM:
            return "ABCDM"
        return cidade if cidade in CORES_REGIAO else "OUTRAS"

    @staticmethod
    def gerar_tipo_servico(df: pd.DataFrame) -> pd.Series:
        if df.empty:
            return pd.Series(dtype="object")

        if CRITERIOS.disponivel and CRITERIOS.classificar:
            try:
                _, serie = CRITERIOS.classificar(df.copy())
                return serie.map(padronizar_tipo_servico)
            except Exception as erro:
                st.warning(f"⚠️ Erro ao usar critérios: {erro}. Usando fallback.")

        col = DF.buscar_coluna(df, Config.COLUNAS_TIPO_SERVICO)
        if col is None:
            return pd.Series("Outros", index=df.index, dtype="object")
        return DF.obter_serie(df, col).map(padronizar_tipo_servico)

    @staticmethod
    def criar_flag_gpon(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            df["FLAG_GPON"] = "Não"
            return df

        if CRITERIOS.disponivel and CRITERIOS.criar_flag_gpon:
            try:
                df_result, _, _ = CRITERIOS.criar_flag_gpon(df)
                return df_result
            except Exception as erro:
                st.warning(f"⚠️ Erro ao criar FLAG_GPON: {erro}. Usando fallback.")

        col = DF.buscar_coluna(
            df,
            ("HABILIDADE DE TRABALHO", "HABILIDADE", "HABILIDADES", "SKILL", "SKILLS"),
        )
        if col:
            mask = (
                df[col].fillna("").astype(str).str.upper().str.contains("PON", na=False)
            )
            df["FLAG_GPON"] = np.where(mask, "Sim", "Não")
        else:
            df["FLAG_GPON"] = "Não"
        return df

    @staticmethod
    def _mesclar_ativos(
        base: pd.DataFrame, ativos: pd.DataFrame, col_login: str | None
    ) -> pd.DataFrame:
        """
        Merge DIRETO e RESILIENTE entre 'LOGIN DO TÉCNICO' (da O.S.) e 'LOGIN' (da lista_ativos).
        Converte o Login da O.S. no NOME DO TÉCNICO e NOME DO MONITOR correspondentes.
        """
        col_mon_os = DF.buscar_coluna(
            base, ("MONITOR", "GESTOR", "SUPERVISOR", "COORDENADOR")
        )
        col_tec_os = DF.buscar_coluna(
            base, ("TÉCNICO", "TECNICO", "NOME TÉCNICO", "NOME TECNICO", "COLABORADOR")
        )

        mon_os = (
            DF.obter_serie(base, col_mon_os).map(normalizar_texto)
            if col_mon_os
            else pd.Series("", index=base.index)
        )
        tec_os = (
            DF.obter_serie(base, col_tec_os).map(normalizar_texto)
            if col_tec_os
            else pd.Series("", index=base.index)
        )

        # Se ativos não estiver disponível, usa os valores da O.S.
        if ativos is None or ativos.empty or "LOGIN" not in ativos.columns:
            base["TÉCNICO"] = tec_os.mask(
                tec_os.isin(VALORES_AUSENTES) | (tec_os == ""), "NÃO MAPEADO"
            )
            base["MONITOR"] = mon_os.mask(
                mon_os.isin(VALORES_AUSENTES) | (mon_os == ""), "SEM MONITOR"
            )
            return base

        # 1. Preparar lista_ativos
        at = ativos.copy()
        at["_JOIN_KEY"] = at["LOGIN"].map(normalizar_login)
        at["_JOIN_CLEAN"] = at["_JOIN_KEY"].str.lstrip("0")

        for col_def in ["TÉCNICO", "MONITOR", "BASE"]:
            if col_def not in at.columns:
                at[col_def] = ""

        at_key = at.drop_duplicates(subset=["_JOIN_KEY"], keep="last")
        at_clean = at.loc[at["_JOIN_CLEAN"] != ""].drop_duplicates(
            subset=["_JOIN_CLEAN"], keep="last"
        )

        # 2. Preparar chaves da O.S.
        os_login = (
            DF.obter_serie(base, col_login)
            if col_login
            else pd.Series("", index=base.index)
        )
        base_work = base.copy()
        base_work["_JOIN_KEY"] = os_login.map(normalizar_login)
        base_work["_JOIN_CLEAN"] = base_work["_JOIN_KEY"].str.lstrip("0")
        base_work["_MON_OS"] = mon_os
        base_work["_TEC_OS"] = tec_os

        cols_rename = {"TÉCNICO": "_TEC_AT", "MONITOR": "_MON_AT", "BASE": "_BASE_AT"}

        # 3. Merge 1: Exato pelo LOGIN
        m1 = base_work.merge(
            at_key[["_JOIN_KEY", "TÉCNICO", "MONITOR", "BASE"]].rename(
                columns=cols_rename
            ),
            on="_JOIN_KEY",
            how="left",
        )

        # 4. Merge 2: Fallback sem zeros à esquerda (ex: 001234 -> 1234)
        m2 = base_work.merge(
            at_clean[["_JOIN_CLEAN", "TÉCNICO", "MONITOR", "BASE"]].rename(
                columns=cols_rename
            ),
            on="_JOIN_CLEAN",
            how="left",
        )

        # 5. Consolidação de Colunas (Ativos Exato -> Ativos Clean -> O.S. Original -> Valor Padrão)

        # Consolidação Técnico (Puxa o NOME DO TÉCNICO de lista_ativos)
        t_at1 = m1["_TEC_AT"].fillna("").astype(str).str.strip()
        t_at2 = m2["_TEC_AT"].fillna("").astype(str).str.strip()
        t_orig = m1["_TEC_OS"].fillna("").astype(str).str.strip()

        t_final = np.where(t_at1 != "", t_at1, np.where(t_at2 != "", t_at2, t_orig))
        t_final = np.where(
            (t_final != "") & (~pd.Series(t_final).isin(VALORES_AUSENTES)),
            t_final,
            "NÃO MAPEADO",
        )

        # Consolidação Monitor (Puxa o NOME DO MONITOR de lista_ativos)
        m_at1 = m1["_MON_AT"].fillna("").astype(str).str.strip()
        m_at2 = m2["_MON_AT"].fillna("").astype(str).str.strip()
        m_orig = m1["_MON_OS"].fillna("").astype(str).str.strip()

        m_final = np.where(m_at1 != "", m_at1, np.where(m_at2 != "", m_at2, m_orig))
        m_final = np.where(
            (m_final != "") & (~pd.Series(m_final).isin(VALORES_AUSENTES)),
            m_final,
            "SEM MONITOR",
        )

        # Consolidação de Base
        b_at1 = m1["_BASE_AT"].fillna("").astype(str).str.strip()
        b_at2 = m2["_BASE_AT"].fillna("").astype(str).str.strip()
        b_final = np.where(b_at1 != "", b_at1, b_at2)

        # Atribuição Garantida mantendo o índice original da base
        res = base.copy()
        res["TÉCNICO"] = DF.limpar_texto(
            pd.Series(t_final, index=res.index), "NÃO MAPEADO"
        )
        res["MONITOR"] = DF.limpar_texto(
            pd.Series(m_final, index=res.index), "SEM MONITOR"
        )
        if any(b_final != ""):
            res["_B"] = pd.Series(b_final, index=res.index)

        return res

    @staticmethod
    def _definir_regiao(base: pd.DataFrame) -> pd.Series:
        col_cidade = DF.buscar_coluna(
            base, ("CIDADE", "LOCALIDADE", "MUNICÍPIO", "MUNICIPIO")
        )
        if col_cidade:
            return DF.obter_serie(base, col_cidade).map(DataLoader.classificar_regiao)
        if "_B" in base.columns:
            return base["_B"].map(DataLoader.classificar_regiao)
        if "BASE" in base.columns:
            return base["BASE"].map(DataLoader.classificar_regiao)
        return pd.Series("OUTRAS", index=base.index)

    @staticmethod
    def preparar_base(
        df: pd.DataFrame,
        df_gs: pd.DataFrame | None,
        filename: str = "",
    ) -> pd.DataFrame:
        if df is None or df.empty:
            return pd.DataFrame()

        base = DF.padronizar_colunas(df)
        total_importado = len(base)

        base = DF.remover_duplicatas_coluna(base)
        base = DF.corrigir_colunas_vazias(base)
        base = DataLoader.criar_flag_gpon(base)
        flag_gpon_sim = int((base["FLAG_GPON"] == "Sim").sum())

        col_tipo_origem = DF.buscar_coluna(base, Config.COLUNAS_TIPO_SERVICO)
        valores_tipo_origem: dict[str, int] = {}
        if col_tipo_origem:
            serie = (
                DF.obter_serie(base, col_tipo_origem).fillna("").astype(str).str.strip()
            )
            valores_tipo_origem = {
                str(k): int(v) for k, v in serie.value_counts().head(20).items()
            }

        base["STATUS CONTRATO"] = ClassificadorStatus.classificar(base)
        status_up = (
            base["STATUS CONTRATO"].fillna("").astype(str).str.upper().str.strip()
        )
        mask_remover = status_up.isin({"CANCELADO", "SUSPENSO"})
        removidos_status = int(mask_remover.sum())
        base = base.loc[~mask_remover].copy()
        if base.empty:
            return pd.DataFrame()

        col_contrato = DF.buscar_coluna(
            base,
            (
                "CONTRATO",
                "NR CONTRATO",
                "NÚMERO CONTRATO",
                "NUMERO CONTRATO",
                "CONTRATO_ID",
            ),
        )
        removidos_contrato = 0
        if col_contrato:
            contratos = DF.obter_serie(base, col_contrato).map(normalizar_texto)
            mask_inv = contratos.isin(
                {"", "NAN", "NONE", "NULL", "N/A", "NA", "-", "0"}
            )
            removidos_contrato = int(mask_inv.sum())
            base = base.loc[~mask_inv].copy()
        if base.empty:
            return pd.DataFrame()

        col_total = DF.buscar_coluna(
            base, ("TOTAL DE TAREFAS", "QTD TAREFAS", "QUANTIDADE", "VOLUME")
        )
        if col_total:
            n = DF.obter_serie(base, col_total).map(converter_numero)
            base["TOTAL DE TAREFAS"] = (
                pd.to_numeric(n, errors="coerce")
                .fillna(1)
                .clip(lower=0)
                .round()
                .astype(int)
            )
        else:
            base["TOTAL DE TAREFAS"] = 1

        # Mapeamento estrito com prioridade máxima para LOGIN DO TÉCNICO
        col_login = DF.buscar_coluna(
            base,
            (
                "LOGIN DO TÉCNICO",
                "LOGIN DO TECNICO",
                "LOGIN_DO_TECNICO",
                "LOGIN TÉCNICO",
                "LOGIN TECNICO",
                "LOGIN_TECNICO",
                "LOGIN",
                "MATRÍCULA",
                "MATRICULA",
                "RE",
                "ID TÉCNICO",
                "ID TECNICO",
                "ID_TECNICO",
                "COD_TECNICO",
            ),
        )

        ativos = (
            DataLoader.processar_lista_ativos(df_gs)
            if isinstance(df_gs, pd.DataFrame)
            else pd.DataFrame()
        )
        base = DataLoader._mesclar_ativos(base, ativos, col_login)

        base["REGIÃO"] = DataLoader._definir_regiao(base)
        base["TIPO_SERVICO"] = DataLoader.gerar_tipo_servico(base).map(
            padronizar_tipo_servico
        )
        base["Status Contrato"] = base["STATUS CONTRATO"]

        aux = [
            c
            for c in base.columns
            if c.startswith("_LOGIN_")
            or c.startswith("_ATIVO")
            or c in ("_T", "_M", "_B")
        ]
        base = base.drop(columns=aux, errors="ignore").reset_index(drop=True)

        base.attrs.update(
            {
                "arquivo_origem": filename,
                "total_importado": total_importado,
                "removidos_suspensos_cancelados": removidos_status,
                "removidos_contrato": removidos_contrato,
                "tipo_servico_coluna": col_tipo_origem or "",
                "valores_tipo_servico_originais": valores_tipo_origem,
                "total_processado": len(base),
                "flag_gpon_sim_count": flag_gpon_sim,
            }
        )

        return base

    @staticmethod
    def callback_robo_etl(df_raw: pd.DataFrame, df_gs: pd.DataFrame) -> pd.DataFrame:
        caminho = st.session_state.get("robo_candidato_path", "")
        nome = Path(str(caminho)).name if caminho else "Arquivo_Robo"
        df_proc = DataLoader.preparar_base(df_raw, df_gs, filename=nome)
        st.session_state.update(
            {
                "_robo_df_pronto": df_proc,
                "_robo_nome_arquivo": nome,
                "_robo_dados_disponiveis": not df_proc.empty,
                "df_memoria_raw": df_raw,
                "df_memoria": df_proc,
                "origem_dados": f"Robô ({nome})",
                "robo_hora_sucesso": datetime.now(),
            }
        )
        return df_proc


# ═════════════════════════════════════════════════════════════════════
# MOTOR ANALÍTICO
# ═════════════════════════════════════════════════════════════════════
class Motor:
    @staticmethod
    def _matriz_logic(df: pd.DataFrame) -> pd.DataFrame:
        obrig = {"MONITOR", "TIPO_SERVICO", "Status Contrato", "TOTAL DE TAREFAS"}
        if df is None or df.empty or not obrig.issubset(df.columns):
            return pd.DataFrame()

        t = df.copy()
        t["MONITOR"] = (
            t["MONITOR"].fillna("SEM MONITOR").astype(str).str.strip().str.upper()
        )
        t["TIPO_SERVICO"] = t["TIPO_SERVICO"].apply(padronizar_tipo_servico)
        t["Status Contrato"] = t["Status Contrato"].fillna("").astype(str).str.strip()
        t["TOTAL DE TAREFAS"] = (
            pd.to_numeric(t["TOTAL DE TAREFAS"], errors="coerce")
            .fillna(0)
            .clip(lower=0)
        )

        validos = t.loc[t["Status Contrato"].isin(Config.STATUS_ORDEM)].copy()
        if validos.empty:
            return pd.DataFrame()

        validos["_EXEC"] = np.where(
            validos["Status Contrato"].eq("Executada"), validos["TOTAL DE TAREFAS"], 0
        )
        validos["_NEX"] = np.where(
            validos["Status Contrato"].eq("Não Executada"),
            validos["TOTAL DE TAREFAS"],
            0,
        )

        ag = (
            validos.groupby(["MONITOR", "TIPO_SERVICO"], dropna=False)
            .agg(executados=("_EXEC", "sum"), nao_executados=("_NEX", "sum"))
            .reset_index()
        )
        ag["DENOMINADOR"] = ag["executados"] + ag["nao_executados"]
        ag["PERCENTUAL"] = div_segura(
            ag["nao_executados"], ag["DENOMINADOR"], padrao=np.nan
        )

        monitores = pd.Index(
            sorted(validos["MONITOR"].dropna().unique()), name="MONITOR"
        )
        pivot = ag.pivot(
            index="MONITOR", columns="TIPO_SERVICO", values="PERCENTUAL"
        ).reindex(monitores)
        for tipo in Config.ORDEM_TIPOS:
            if tipo not in pivot.columns:
                pivot[tipo] = np.nan
        pivot = pivot.loc[:, list(Config.ORDEM_TIPOS)]

        totais = (
            validos.groupby("MONITOR")[["_EXEC", "_NEX"]]
            .sum()
            .reindex(monitores)
            .fillna(0)
        )
        denom = totais["_EXEC"] + totais["_NEX"]
        pivot["Quebra Geral"] = div_segura(totais["_NEX"], denom, padrao=np.nan)
        pivot["Total Tasks"] = (
            validos.groupby("MONITOR")["TOTAL DE TAREFAS"]
            .sum()
            .reindex(monitores)
            .fillna(0)
            .round()
            .astype(int)
        )
        pivot = pivot.reset_index().rename(columns={"MONITOR": "Monitor"})

        exec_g = float(validos["_EXEC"].sum())
        nex_g = float(validos["_NEX"].sum())
        denom_g = exec_g + nex_g
        total_row: dict[str, Any] = {
            "Monitor": "TOTAL GERAL",
            "Quebra Geral": nex_g / denom_g if denom_g > 0 else np.nan,
            "Total Tasks": int(validos["TOTAL DE TAREFAS"].sum()),
        }
        for tipo in Config.ORDEM_TIPOS:
            r = validos.loc[validos["TIPO_SERVICO"].eq(tipo)]
            e, n = float(r["_EXEC"].sum()), float(r["_NEX"].sum())
            total_row[tipo] = n / (e + n) if (e + n) > 0 else np.nan

        return pd.concat(
            [pivot, pd.DataFrame([total_row], columns=pivot.columns)], ignore_index=True
        )

    @staticmethod
    def matriz_resumo(df: pd.DataFrame) -> pd.DataFrame:
        if df is None or df.empty:
            return pd.DataFrame()
        return _cache_matriz(df.copy())

    @staticmethod
    def _pivot_status(df: pd.DataFrame, grupo: str) -> pd.DataFrame:
        t = df.copy()
        t["TOTAL DE TAREFAS"] = (
            pd.to_numeric(t["TOTAL DE TAREFAS"], errors="coerce")
            .fillna(0)
            .clip(lower=0)
        )
        pivot = pd.pivot_table(
            t,
            index=grupo,
            columns="Status Contrato",
            values="TOTAL DE TAREFAS",
            aggfunc="sum",
            fill_value=0,
        )
        for s in Config.STATUS_ORDEM:
            if s not in pivot.columns:
                pivot[s] = 0.0
        return pivot.reset_index()

    @staticmethod
    def tabela_cenarios(
        df: pd.DataFrame,
        grupo: str,
        p_ot: float,
        p_base: float,
        p_pess: float,
        min_aloc: float = 5,
    ) -> pd.DataFrame:
        obrig = {grupo, "Status Contrato", "TOTAL DE TAREFAS"}
        if df.empty or not obrig.issubset(df.columns):
            return pd.DataFrame()

        p_ot, p_base, p_pess = (float(np.clip(x, 0, 1)) for x in (p_ot, p_base, p_pess))
        r = Motor._pivot_status(df, grupo)
        r["Considerado"] = r["Executada"] + r["Não Executada"]
        r["Alocado"] = r["Considerado"] + r["Pendente"]
        r["Quebra Atual"] = div_segura(r["Não Executada"], r["Considerado"])

        for nome, prob in (
            ("Otimista", p_ot),
            ("Base", p_base),
            ("Pessimista", p_pess),
        ):
            r[f"Fechamento {nome}"] = div_segura(
                r["Não Executada"] + r["Pendente"] * prob, r["Alocado"]
            )

        return (
            r.loc[r["Alocado"] >= min_aloc]
            .sort_values("Fechamento Base", ascending=False)
            .reset_index(drop=True)
        )

    @staticmethod
    def tecnicos_criticos(
        df: pd.DataFrame,
        segmento: str,
        p_base: float,
        min_aloc: float,
        top_n: int,
        p_ot: float = 0.15,
        p_pess: float = 0.50,
    ) -> pd.DataFrame:
        recorte = (
            df.loc[df["TIPO_SERVICO"].eq(segmento)].copy()
            if segmento and segmento != "TODOS" and "TIPO_SERVICO" in df.columns
            else df.copy()
        )
        if recorte.empty:
            return pd.DataFrame()
        return Motor.tabela_cenarios(
            recorte, "TÉCNICO", p_ot, p_base, p_pess, min_aloc
        ).head(top_n)

    @staticmethod
    def causa_raiz(df: pd.DataFrame, coluna_baixa: str, top_n: int = 8) -> pd.DataFrame:
        if not {"Status Contrato", "TOTAL DE TAREFAS", coluna_baixa}.issubset(
            df.columns
        ):
            return pd.DataFrame()
        r = df.loc[df["Status Contrato"].eq("Não Executada")].copy()
        if r.empty:
            return pd.DataFrame()

        r["_B"] = (
            r[coluna_baixa]
            .fillna("SEM REGISTRO")
            .astype(str)
            .map(normalizar_texto)
            .replace("", "SEM REGISTRO")
        )
        r["TOTAL DE TAREFAS"] = (
            pd.to_numeric(r["TOTAL DE TAREFAS"], errors="coerce")
            .fillna(0)
            .clip(lower=0)
        )

        total = float(r["TOTAL DE TAREFAS"].sum())
        res = r.groupby("_B")["TOTAL DE TAREFAS"].sum().nlargest(top_n).reset_index()
        res.columns = ["Motivo de Baixa", "Volume"]
        res["% do Total"] = div_segura(res["Volume"], pd.Series(total, index=res.index))
        res["Acumulado"] = res["% do Total"].cumsum()
        return res

    @staticmethod
    def backoffice_fila(df: pd.DataFrame) -> pd.DataFrame:
        if df is None or df.empty or "Status Contrato" not in df.columns:
            return pd.DataFrame()

        t = df.copy()
        for col, pad in (
            ("MONITOR", "SEM MONITOR"),
            ("TÉCNICO", "NÃO MAPEADO"),
            ("TIPO_SERVICO", "Outros"),
        ):
            if col not in t.columns:
                t[col] = pad
        if "TOTAL DE TAREFAS" in t.columns:
            t["TOTAL DE TAREFAS"] = (
                pd.to_numeric(t["TOTAL DE TAREFAS"], errors="coerce")
                .fillna(0)
                .clip(lower=0)
            )
        else:
            t["TOTAL DE TAREFAS"] = 0

        fila = t.loc[t["Status Contrato"].isin(["Não Executada", "Pendente"])].copy()
        if fila.empty:
            return pd.DataFrame()

        ag = (
            fila.groupby(
                ["MONITOR", "TÉCNICO", "TIPO_SERVICO", "Status Contrato"], dropna=False
            )["TOTAL DE TAREFAS"]
            .sum()
            .reset_index()
        )
        pivot = pd.pivot_table(
            ag,
            index=["MONITOR", "TÉCNICO", "TIPO_SERVICO"],
            columns="Status Contrato",
            values="TOTAL DE TAREFAS",
            aggfunc="sum",
            fill_value=0,
        ).reset_index()

        for c in ("Não Executada", "Pendente"):
            if c not in pivot.columns:
                pivot[c] = 0

        pivot["Total Fila"] = pivot["Não Executada"] + pivot["Pendente"]
        pivot["Prioridade"] = pivot["Não Executada"] * 2 + pivot["Pendente"]
        pivot["Classificação"] = np.select(
            [
                pivot["Prioridade"].ge(20),
                pivot["Prioridade"].ge(10),
                pivot["Prioridade"].ge(5),
            ],
            ["🔴 CRÍTICO", "🟠 ALTA", "🟡 MÉDIA"],
            default="⚪ BAIXA",
        )

        pivot = pivot.rename(
            columns={
                "MONITOR": "Monitor",
                "TÉCNICO": "Técnico",
                "TIPO_SERVICO": "Segmento",
            }
        )
        return pivot.sort_values("Prioridade", ascending=False).reset_index(drop=True)[
            [
                "Classificação",
                "Monitor",
                "Técnico",
                "Segmento",
                "Não Executada",
                "Pendente",
                "Total Fila",
                "Prioridade",
            ]
        ]

    @staticmethod
    def projetar(
        df: pd.DataFrame,
        probabilidade: float = 0.30,
        grupo: str = "MONITOR",
        incluir_meta: bool = True,
        meta_sla: float = Config.SLA_QUEBRA_MAXIMA,
    ) -> pd.DataFrame:
        obrig = {grupo, "Status Contrato", "TOTAL DE TAREFAS"}
        if not obrig.issubset(df.columns):
            return pd.DataFrame()

        probabilidade = float(np.clip(probabilidade, 0, 1))
        meta_sla = float(np.clip(meta_sla, 0, 1))

        r = Motor._pivot_status(df, grupo)
        r["Considerado"] = r["Executada"] + r["Não Executada"]
        r["Alocado"] = r["Considerado"] + r["Pendente"]
        r["Quebra Atual"] = div_segura(r["Não Executada"], r["Considerado"])
        r["Projeção Fechamento"] = div_segura(
            r["Não Executada"] + r["Pendente"] * (1 - probabilidade),
            r["Alocado"],
        )
        r["Executadas Projetadas"] = r["Executada"] + r["Pendente"] * probabilidade

        num = r["Alocado"] * (1 - meta_sla) - r["Executada"]
        conv = div_segura(num, r["Pendente"]).clip(0, 1)
        r["Conversão Necessária"] = np.where(r["Pendente"].gt(0), conv, 0.0)

        if incluir_meta:
            p = r["Projeção Fechamento"]
            r["Status Meta"] = np.select(
                [p.le(meta_sla), p.le(meta_sla * 1.25), p.le(meta_sla * 1.50)],
                ["✅ Dentro da Meta", "⚠️ Atenção", "🔶 Crítico"],
                default="🚨 Muito Crítico",
            )

        for c in ("Executada", "Não Executada", "Pendente", "Considerado", "Alocado"):
            r[c] = r[c].fillna(0).round().astype(int)

        return r.sort_values("Projeção Fechamento", ascending=False).reset_index(
            drop=True
        )

    @staticmethod
    def resumo_gpon(df: pd.DataFrame) -> dict[str, Any]:
        if df is None or df.empty or "FLAG_GPON" not in df.columns:
            return {"total": 0, "sim": 0, "nao": 0, "percentual": 0.0}
        total, sim = len(df), int((df["FLAG_GPON"] == "Sim").sum())
        return {
            "total": total,
            "sim": sim,
            "nao": total - sim,
            "percentual": (sim / total * 100) if total > 0 else 0.0,
        }


@st.cache_data(ttl=3600, show_spinner=False)
def _cache_matriz(df: pd.DataFrame) -> pd.DataFrame:
    return Motor._matriz_logic(df)


# ═════════════════════════════════════════════════════════════════════
# CAMADA VISUAL
# ═════════════════════════════════════════════════════════════════════
_TEMAS_KPI: Final[dict[str, dict[str, str]]] = {
    "amarelo": {
        "fundo": "#FEF9C3",
        "texto": "#854D0E",
        "borda": "#EAB308",
        "titulo": "#A16207",
    },
    "roxo": {
        "fundo": "#FAF5FF",
        "texto": "#7E22CE",
        "borda": "#A855F7",
        "titulo": "#6B21A8",
    },
    "escuro": {
        "fundo": "#1E293B",
        "texto": "#FFFFFF",
        "borda": "#475569",
        "titulo": "#E2E8F0",
    },
}


def render_section_header(icone: str, titulo: str) -> None:
    if COMPONENTES.section_header:
        COMPONENTES.section_header(icone, titulo)
    else:
        st.subheader(f"{icone} {titulo}" if icone else titulo)


def render_sidebar_brand(**kw: Any) -> None:
    if COMPONENTES.sidebar_brand:
        COMPONENTES.sidebar_brand(**kw)


def render_table_html(df: pd.DataFrame, **kw: Any) -> None:
    if COMPONENTES.table_html:
        COMPONENTES.table_html(df, **kw)
    else:
        st.dataframe(df, use_container_width=True, hide_index=True)


def render_insight(texto: str, tipo: TipoInsight = "info") -> None:
    if COMPONENTES.insight:
        COMPONENTES.insight(texto, tipo)
        return
    if tipo == "ok":
        st.success(texto)
    elif tipo in {"critico", "alerta"}:
        st.warning(texto)
    elif tipo == "acao":
        st.error(texto)
    else:
        st.info(texto)


def render_kpi_sm(
    col: Any, label: str, value: str, sub: str = "", tema: TemaKPI = "azul"
) -> None:
    if tema not in _TEMAS_KPI:
        if COMPONENTES.kpi_sm:
            COMPONENTES.kpi_sm(col, label, value, sub, tema)
        else:
            col.metric(label, value, sub)
        return

    e = _TEMAS_KPI[tema]
    col.markdown(
        f"""
        <div style="background:{e['fundo']};border-left:3px solid {e['borda']};border-radius:6px;
                    padding:12px 16px;margin-bottom:8px;box-shadow:0 1px 4px rgba(0,0,0,0.06);">
            <div style="font-family:{Fontes.TEXTO};font-size:10px;color:{e['titulo']};
                        text-transform:uppercase;letter-spacing:1px;font-weight:700;">{html_safe(label)}</div>
            <div style="font-family:{Fontes.TITULO};font-size:20px;color:{e['texto']};
                        font-weight:800;line-height:1.2;margin-top:4px;">{html_safe(value)}</div>
            <div style="font-family:{Fontes.TEXTO};font-size:11px;color:{e['titulo']};margin-top:2px;">{html_safe(sub)}</div>
        </div>
    """,
        unsafe_allow_html=True,
    )


def renderizar_robo_local(*args: Any, **kw: Any) -> None:
    if _ROBO_FN is None:
        st.sidebar.warning("🤖 Robô offline. Utilize upload manual.")
        return
    try:
        _ROBO_FN(*args, **kw)
    except Exception as erro:
        st.sidebar.error(f"Erro no Robô: {erro}")


# ═════════════════════════════════════════════════════════════════════
# CSS GLOBAL
# ═════════════════════════════════════════════════════════════════════
_CSS_GLOBAL: Final[str] = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Manrope:wght@400;600;700;800&display=swap');
.import-header-title{font-family:'Manrope',sans-serif;font-size:32px;font-weight:800;color:#1E293B;letter-spacing:-0.5px;margin:0;display:inline-block}
.import-header-badge{display:inline-flex;align-items:center;background:#EEF2FF;color:#3B82F6;border:1.5px solid #93C5FD;font-size:11px;font-weight:800;padding:3px 12px;border-radius:9999px;letter-spacing:0.6px;text-transform:uppercase;margin-left:14px;vertical-align:middle}
.import-header-sub{font-size:14px;color:#64748B;margin-top:6px;margin-bottom:14px;font-family:'Inter',sans-serif}
.import-header-line{height:3px;background:linear-gradient(90deg,#012869 0%,#1E40AF 30%,#F59E0B 65%,#EF4444 100%);border-radius:2px;margin-bottom:18px}
.import-alert-green{background-color:#E8F8F0;border:1px solid #C2F0D9;border-radius:8px;padding:14px 20px;color:#107C41;font-size:14px;font-weight:700;display:flex;align-items:center;gap:10px;margin-bottom:16px;font-family:'Inter',sans-serif}
.hero-container{background:rgba(248,250,252,0.95);padding:0.5rem 0;border-radius:14px;margin-bottom:20px}
.hero-card{background:linear-gradient(135deg,#012869 0%,#1E40AF 50%,#F37C04 100%);padding:28px 40px;border-radius:14px;color:white;box-shadow:0 10px 40px rgba(1,40,105,0.20)}
.hero-title{margin:0;font-size:30px;font-weight:800;color:white!important;font-family:'Manrope',sans-serif}
.hero-sub{margin:6px 0 0 0;font-size:14px;color:#F8FAFC;font-family:'Inter',sans-serif}
.base-info{background:linear-gradient(135deg,#0F172A 0%,#1E3A5F 100%);padding:1rem 1.5rem;border-radius:0.75rem;margin-bottom:1.5rem;display:flex;align-items:center;flex-wrap:wrap;gap:0.6rem;box-shadow:0 4px 12px rgba(0,0,0,0.15)}
.badge-regiao{padding:0.3rem 0.9rem;border-radius:999px;font-size:0.82rem;font-weight:700;border:2px solid}
.matriz-table-container{width:100%;overflow-x:auto;border-radius:12px;box-shadow:0 4px 16px rgba(0,0,0,0.06);border:1px solid #E2E8F0;background:#FFFFFF;margin-bottom:16px;font-family:'Inter',sans-serif}
.matriz-table{width:100%;border-collapse:collapse;font-size:13px}
.matriz-table th{background:#F8FAFC;color:#475569;font-weight:700;text-transform:uppercase;letter-spacing:0.05em;padding:14px 16px;border-bottom:2px solid #E2E8F0;text-align:center}
.matriz-table th:first-child{text-align:left}
.matriz-table td{padding:10px 16px;border-bottom:1px solid #F1F5F9;text-align:center;color:#1E293B}
.matriz-table td:first-child{text-align:left;font-weight:600}
.matriz-table tr:hover{background-color:#F8FAFC}
.matriz-table tr.total-row{background-color:#F1F5F9;font-weight:800;border-top:2px solid #CBD5E1}
.badge-meta-ok{background-color:#DCFCE7;color:#15803D;font-weight:700;padding:5px 12px;border-radius:6px;display:inline-block;min-width:70px;border:1px solid #86EFAC}
.badge-meta-nok{background-color:#FEE2E2;color:#B91C1C;font-weight:700;padding:5px 12px;border-radius:6px;display:inline-block;min-width:70px;border:1px solid #FCA5A5}
.badge-gpon-sim{background-color:#DBEAFE;color:#1E40AF;font-weight:700;padding:4px 10px;border-radius:4px;display:inline-block;font-size:11px}
.badge-gpon-nao{background-color:#F1F5F9;color:#475569;font-weight:600;padding:4px 10px;border-radius:4px;display:inline-block;font-size:11px}
</style>
"""


def injetar_css() -> None:
    st.markdown(_CSS_GLOBAL, unsafe_allow_html=True)


# ═════════════════════════════════════════════════════════════════════
# RENDERS COMPLEXOS
# ═════════════════════════════════════════════════════════════════════
_METAS_COLUNA_MATRIZ: Final[dict[str, str]] = {
    "NOVOS DOMICÍLIOS": "geral",
    "NOVOS DOMICILIOS": "geral",
    "PME": "geral",
    "MIGRAÇÃO": "migracao",
    "MIGRACAO": "migracao",
    "OUTROS": "geral",
    "QUEBRA GERAL": "geral",
}


def render_matriz_executiva(
    df: pd.DataFrame, meta_geral: float = Config.SLA_QUEBRA_MAXIMA
) -> None:
    if df.empty:
        st.info("Sem dados na Matriz Executiva.")
        return

    metas = {
        k: (meta_geral if v == "geral" else Config.SLA_MIGRACAO_MAXIMA)
        for k, v in _METAS_COLUNA_MATRIZ.items()
    }

    linhas: list[str] = [
        '<div class="matriz-table-container"><table class="matriz-table"><thead><tr>'
    ]
    linhas.extend(f"<th>{html_safe(c)}</th>" for c in df.columns)
    linhas.append("</tr></thead><tbody>")

    for _, linha in df.iterrows():
        is_total = str(linha.iloc[0]).strip().upper() == "TOTAL GERAL"
        linhas.append("<tr class='total-row'>" if is_total else "<tr>")

        for i, col in enumerate(df.columns):
            v = linha.iloc[i]
            cu = str(col).strip().upper()

            if cu in metas:
                try:
                    n = float(v)
                    if not np.isfinite(n):
                        linhas.append(
                            "<td><span style='color:#94A3B8;font-weight:700;'>—</span></td>"
                        )
                    else:
                        cls = "badge-meta-ok" if n <= metas[cu] else "badge-meta-nok"
                        linhas.append(
                            f"<td><span class='{cls}'>{fmt_pct_br(n)}</span></td>"
                        )
                except (TypeError, ValueError):
                    linhas.append(f"<td>{html_safe(v)}</td>")
            elif cu in {"TOTAL TASKS", "TOTAL_TASKS"}:
                linhas.append(f"<td><strong>{fmt_int_br(v)}</strong></td>")
            else:
                conteudo = (
                    html_safe(v)
                    if not is_missing(v)
                    else "<span style='color:#94A3B8;'>—</span>"
                )
                linhas.append(f"<td>{conteudo}</td>")
        linhas.append("</tr>")
    linhas.append("</tbody></table></div>")

    st.markdown("".join(linhas), unsafe_allow_html=True)


def render_bloco_importacao(dados_prontos: bool = False) -> bool:
    st.markdown(
        """
        <div style="margin-top:6px;margin-bottom:4px;">
            <span class="import-header-title">Importação de Dados</span>
            <span class="import-header-badge">EXCEL, CSV OU AUTO</span>
        </div>
        <div class="import-header-sub">Envie a base consolidada de O.S. do dia ou aguarde o robô detectar automaticamente.</div>
        <div class="import-header-line"></div>
    """,
        unsafe_allow_html=True,
    )

    if not dados_prontos:
        return False

    st.markdown(
        """
        <div class="import-alert-green">
            <span style="font-size:16px;">✅</span>
            <span>Dados carregados e processados. O relatório está pronto.</span>
        </div>
    """,
        unsafe_allow_html=True,
    )

    return st.button(
        "🔄 Atualizar Relatório",
        type="primary",
        use_container_width=True,
        key="btn_atualizar_relatorio",
    )


def _html_base_ativa(regioes: list[str], total: int, origem: str) -> str:
    badges = []
    for regiao in sorted({str(r).upper().strip() for r in regioes}):
        c = CORES_REGIAO.get(regiao, CORES_REGIAO["OUTRAS"])
        badges.append(
            f'<span class="badge-regiao" style="background:{c["bg"]};color:{c["text"]};border-color:{c["border"]};">{html_safe(regiao)}</span>'
        )

    origem_html = (
        f'<span style="color:#6EE7B7;font-size:0.78rem;font-weight:600;margin-left:8px;">• {html_safe(origem)}</span>'
        if origem
        else ""
    )

    return (
        '<div class="base-info">'
        '<span style="color:#94A3B8;font-size:0.8rem;font-weight:700;text-transform:uppercase;letter-spacing:0.08em;">📋 Base Ativa:</span>'
        + "".join(badges)
        + origem_html
        + f'<span style="color:#FFFFFF;font-size:0.78rem;margin-left:auto;font-weight:700;">{fmt_int_br(total)} registros</span>'
        "</div>"
    )


def render_hero(
    titulo: str,
    subtitulo: str,
    regioes: list[str],
    total: int,
    badge: str = "",
    origem: str = "",
) -> None:
    badge_html = (
        f'<span style="display:inline-block;background:rgba(255,255,255,0.20);padding:5px 16px;'
        f"border-radius:20px;font-size:12px;font-weight:700;margin-top:10px;letter-spacing:0.6px;"
        f'text-transform:uppercase;color:white;border:1px solid rgba(255,255,255,0.30);">{html_safe(badge)}</span>'
        if badge
        else ""
    )
    base_html = _html_base_ativa(regioes, total, origem) if total > 0 else ""

    st.markdown(
        f'<div class="hero-container"><div class="hero-card">'
        f'<h1 class="hero-title">{html_safe(titulo)}</h1>'
        f'<p class="hero-sub">{html_safe(subtitulo)}</p>'
        f"{badge_html}</div>{base_html}</div>",
        unsafe_allow_html=True,
    )


_PALAVRAS_PCT_TABELA: Final[tuple[str, ...]] = (
    "QUEBRA",
    "FECHAMENTO",
    "%",
    "PROJEÇÃO",
    "PROJECAO",
    "CONVERSÃO",
    "CONVERSAO",
    "ACUMULADO",
)
_PALAVRAS_INT_TABELA: Final[tuple[str, ...]] = (
    "TOTAL",
    "VOLUME",
    "TAREFAS",
    "PRIORIDADE",
    "EXECUTADA",
    "TASKS",
    "ALOCADO",
    "CONSIDERADO",
)


def render_tabela(
    df: pd.DataFrame,
    titulo: str,
    icone: str,
    color_col: str | None = None,
    meta: float = Config.SLA_QUEBRA_MAXIMA,
    height: int = 400,
) -> None:
    st.markdown(
        f"""
        <div style="background:#FFFFFF;border-radius:0.75rem;padding:1rem 1.2rem;
                    box-shadow:0 2px 8px rgba(0,0,0,0.05);margin-bottom:0.5rem;">
            <div style="font-size:1rem;font-weight:700;color:#0F172A;display:flex;
                        align-items:center;gap:0.5rem;">
                <span>{html_safe(icone)}</span>
                <span>{html_safe(titulo)}</span>
                <span style="font-size:0.68rem;background:#E0F2FE;color:#0369A1;
                             padding:0.15rem 0.5rem;border-radius:999px;">{len(df)} registros</span>
            </div>
        </div>
    """,
        unsafe_allow_html=True,
    )

    if df is None or df.empty:
        st.info("Sem dados para exibir.")
        return

    formatos: dict[str, str] = {}
    for col in df.columns:
        cu = str(col).upper()
        if any(k in cu for k in _PALAVRAS_PCT_TABELA):
            formatos[str(col)] = "{:.2%}"
        elif any(k in cu for k in _PALAVRAS_INT_TABELA):
            formatos[str(col)] = "{:,.0f}"

    regras = None
    if color_col and color_col in df.columns:
        regras = {
            "coluna": color_col,
            "meta": meta,
            "acima_meta": {"bg": "#FEE2E2", "text": "#991B1B", "bold": True},
            "abaixo_meta": {"bg": "#DCFCE7", "text": "#166534", "bold": True},
        }

    render_table_html(
        df,
        fmt=formatos,
        color_rules=regras,
        colunas_num=df.select_dtypes(include=np.number).columns.astype(str).tolist(),
        height=height,
    )


def render_card_gpon(df: pd.DataFrame) -> None:
    r = Motor.resumo_gpon(df)
    st.markdown(
        f"""
        <div style="background:linear-gradient(135deg,#0369A1 0%,#0284C7 100%);
                    padding:16px 20px;border-radius:10px;color:white;margin-bottom:16px;">
            <div style="display:flex;justify-content:space-between;align-items:center;">
                <div>
                    <div style="font-size:11px;text-transform:uppercase;letter-spacing:0.6px;opacity:0.9;">📡 FLAG_GPON</div>
                    <div style="font-size:24px;font-weight:800;margin-top:4px;">{r['sim']:,}</div>
                    <div style="font-size:11px;opacity:0.85;margin-top:2px;">{r['percentual']:.1f}% com GPON habilitado</div>
                </div>
                <div style="text-align:right;">
                    <div style="font-size:11px;opacity:0.85;">Total</div>
                    <div style="font-size:18px;font-weight:700;">{r['total']:,}</div>
                </div>
            </div>
        </div>
    """,
        unsafe_allow_html=True,
    )


# ═════════════════════════════════════════════════════════════════════
# VIEWS
# ═════════════════════════════════════════════════════════════════════
def view_resumo_executivo(df: pd.DataFrame, meta_sla: float) -> None:
    col_tipo = str(df.attrs.get("tipo_servico_coluna", ""))
    tipos_proc = set(
        df.get("TIPO_SERVICO", pd.Series(dtype="object")).dropna().astype(str).unique()
    )

    if not col_tipo:
        st.warning(
            "⚠️ A base não possui uma coluna identificável de tipo de serviço. "
            "Todos os registros foram classificados como 'Outros'."
        )

    if tipos_proc and tipos_proc <= {"Outros"}:
        originais = df.attrs.get("valores_tipo_servico_originais", {})
        st.warning(
            "⚠️ A coluna de segmento foi encontrada, mas nenhum valor foi reconhecido "
            "como Novos Domicílios, PME ou Migração."
        )
        if isinstance(originais, dict) and originais:
            st.caption(
                "Valores originais encontrados: "
                + ", ".join(list(originais.keys())[:10])
            )

    if "FLAG_GPON" in df.columns:
        render_card_gpon(df)

    matriz = Motor.matriz_resumo(df)
    if matriz.empty:
        st.warning("⚠️ Dados insuficientes para montar a Matriz Executiva.")
        return

    if "Monitor" in matriz.columns:
        mask = matriz["Monitor"].astype(str).str.strip().str.upper().eq("TOTAL GERAL")
        matriz = pd.concat([matriz.loc[~mask], matriz.loc[mask]], ignore_index=True)

    render_matriz_executiva(matriz, meta_geral=meta_sla)
    st.download_button(
        "📥 Baixar Matriz (Excel)",
        data=gerar_excel(matriz, "Matriz_Resumo"),
        file_name="Matriz_Resumo_Quebra.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True,
    )


def view_analise_detalhada(
    df: pd.DataFrame,
    p_ot: float,
    p_base: float,
    p_pess: float,
    min_aloc: float,
    meta_sla: float,
) -> None:
    if CRITERIOS.disponivel and CRITERIOS.render_painel:
        with st.expander("📋 Painel de Critérios de Classificação", expanded=False):
            CRITERIOS.render_painel(df)

    tab1, tab2, tab3, tab4, tab5 = st.tabs(
        [
            "📉 Projeções",
            "🎯 Projeção Customizada",
            "🏆 Rankings",
            "🔍 Causas Raiz",
            "📋 Backoffice",
        ]
    )

    with tab1:
        render_section_header("📈", "Cenários de Projeção de Fechamento")
        render_tabela(
            Motor.tabela_cenarios(df, "MONITOR", p_ot, p_base, p_pess, min_aloc),
            "Projeção por Monitor",
            "👨‍💼",
            color_col="Fechamento Base",
            meta=meta_sla,
        )

    with tab2:
        render_section_header("🎯", "Projeção Customizada")
        c1, c2 = st.columns(2)
        with c1:
            prob = st.slider("Probabilidade de Conversão (%)", 0, 100, 30, step=5) / 100
        with c2:
            grupo = st.selectbox("Agrupar por", ["MONITOR", "TÉCNICO", "REGIÃO"])

        proj = Motor.projetar(df, probabilidade=prob, grupo=grupo, meta_sla=meta_sla)
        render_tabela(
            proj,
            f"Projeção por {grupo}",
            "🎯",
            color_col="Projeção Fechamento",
            meta=meta_sla,
        )

        if not proj.empty:
            k1, k2, k3, k4 = st.columns(4)
            total = len(proj)
            dentro = int(proj["Projeção Fechamento"].le(meta_sla).sum())
            atencao = int(
                (
                    proj["Projeção Fechamento"].gt(meta_sla)
                    & proj["Projeção Fechamento"].le(meta_sla * 1.25)
                ).sum()
            )
            critico = int(proj["Projeção Fechamento"].gt(meta_sla * 1.25).sum())
            conv_med = float(proj["Conversão Necessária"].mean())

            def _pct(x: int) -> str:
                return f"{x / total:.0%}" if total > 0 else "0%"

            render_kpi_sm(k1, "Dentro da Meta", str(dentro), _pct(dentro), "verde")
            render_kpi_sm(k2, "Atenção", str(atencao), _pct(atencao), "amarelo")
            render_kpi_sm(k3, "Crítico", str(critico), _pct(critico), "vermelho")
            render_kpi_sm(
                k4, "Conversão Média", f"{conv_med:.1%}", "Para atingir meta", "azul"
            )

    with tab3:
        render_section_header("🏆", "Técnicos Mais Críticos")
        render_tabela(
            Motor.tecnicos_criticos(df, "TODOS", p_base, min_aloc, 20, p_ot, p_pess),
            "Top 20 Técnicos com Maior Risco",
            "👤",
            color_col="Fechamento Base",
            meta=meta_sla,
        )

    with tab4:
        render_section_header("🔍", "Análise de Causa Raiz")
        col = DF.buscar_coluna(
            df, ("CÓD DE BAIXA 1", "COD DE BAIXA 1", "MOTIVO DE BAIXA", "COD_BAIXA")
        )
        if col:
            render_tabela(Motor.causa_raiz(df, col, 10), "Pareto de Motivos", "🎯")
        else:
            st.warning("⚠️ Coluna de Código/Motivo de Baixa não encontrada.")

    with tab5:
        render_section_header("📋", "Gestão de Fila (Backoffice)")
        render_tabela(Motor.backoffice_fila(df), "Fila Priorizada", "📋")


def view_auditoria(df: pd.DataFrame) -> None:
    render_section_header("🔎", "Auditoria de Processamento")

    total = len(df)
    importado = int(df.attrs.get("total_importado", total))
    rem_status = int(df.attrs.get("removidos_suspensos_cancelados", 0))
    rem_contrato = int(df.attrs.get("removidos_contrato", 0))
    gpon_sim = int(df.attrs.get("flag_gpon_sim_count", 0))

    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Importados", fmt_int_br(importado))
    k2.metric("Processados", fmt_int_br(total))
    k3.metric("Removidos por Status", fmt_int_br(rem_status))
    k4.metric("Contratos Inválidos", fmt_int_br(rem_contrato))

    st.caption(
        f"Arquivo de origem: {df.attrs.get('arquivo_origem', 'Não identificado')}"
    )
    st.caption(
        f"Coluna de tipo de serviço: {df.attrs.get('tipo_servico_coluna', 'Não identificada')}"
    )

    if "FLAG_GPON" in df.columns:
        st.markdown("#### 📡 FLAG_GPON")
        g1, g2 = st.columns(2)
        pct_sim = (gpon_sim / total * 100) if total > 0 else 0.0
        pct_nao = ((total - gpon_sim) / total * 100) if total > 0 else 0.0
        g1.metric("FLAG_GPON = Sim", f"{gpon_sim:,}", f"{pct_sim:.1f}%")
        g2.metric("FLAG_GPON = Não", f"{total - gpon_sim:,}", f"{pct_nao:.1f}%")

    if "TIPO_SERVICO" in df.columns:
        st.markdown("### 📊 Distribuição de TIPO_SERVICO")
        dist = df["TIPO_SERVICO"].value_counts().reset_index()
        dist.columns = ["Segmento", "Registros"]
        dist["%"] = (dist["Registros"] / total * 100).round(1).astype(str) + "%"
        st.dataframe(dist, hide_index=True, use_container_width=True)

        nd = int(dist.loc[dist["Segmento"] == "Novos Domicílios", "Registros"].sum())
        pme = int(dist.loc[dist["Segmento"] == "PME", "Registros"].sum())

        if nd == 0 and pme == 0:
            st.error(
                "**⚠️ ALERTA CRÍTICO:** Novos Domicílios e PME estão zerados!\n\n"
                "**Possíveis causas:**\n"
                "1. Coluna TIPO O.S 1 não está sendo detectada\n"
                "2. Termos de busca (ADESAO, INSTALACAO) não correspondem aos dados\n"
                "3. Coluna tem nome diferente do esperado"
            )

    if (
        CRITERIOS.disponivel
        and CRITERIOS.detect_tipo_os
        and CRITERIOS.detect_habilidade
        and CRITERIOS.detect_flag_gpon
    ):
        with st.expander("🔍 Diagnóstico de Critérios", expanded=True):
            col_tipo = CRITERIOS.detect_tipo_os(df)
            col_hab = CRITERIOS.detect_habilidade(df)
            col_gpon = CRITERIOS.detect_flag_gpon(df)

            c1, c2, c3 = st.columns(3)
            c1.metric(
                "TIPO O.S 1",
                "✅ Encontrada" if col_tipo else "❌ Não encontrada",
                col_tipo or "—",
            )
            c2.metric(
                "HABILIDADE",
                "✅ Encontrada" if col_hab else "❌ Não encontrada",
                col_hab or "—",
            )
            c3.metric(
                "FLAG_GPON",
                "✅ Encontrada" if col_gpon else "❌ Não encontrada",
                col_gpon or "—",
            )

            if col_tipo:
                st.markdown("**📋 Amostras de TIPO O.S 1:**")
                st.code(df[col_tipo].dropna().astype(str).unique()[:15].tolist())
            if col_hab:
                st.markdown("**📋 Amostras de HABILIDADE:**")
                st.code(df[col_hab].dropna().astype(str).unique()[:15].tolist())

    with st.expander("Visualizar amostra da base processada", expanded=False):
        st.dataframe(df.head(200), use_container_width=True, hide_index=True)


# ═════════════════════════════════════════════════════════════════════
# GESTÃO DE SESSÃO
# ═════════════════════════════════════════════════════════════════════
_CHAVES_SESSAO: Final[tuple[str, ...]] = (
    "df_memoria",
    "df_memoria_raw",
    "arquivo_processado",
    "origem_dados",
    "_robo_df_pronto",
    "_robo_nome_arquivo",
    "_robo_dados_disponiveis",
    "robo_hora_sucesso",
)
_CHAVES_ATIVOS: Final[tuple[str, ...]] = ("df_gs_manual", "arquivo_ativos_processado")


def limpar_sessao(incluir_ativos: bool = True) -> None:
    chaves = list(_CHAVES_SESSAO) + (list(_CHAVES_ATIVOS) if incluir_ativos else [])
    for k in chaves:
        st.session_state.pop(k, None)
    st.session_state["upload_nonce"] = int(st.session_state.get("upload_nonce", 0)) + 1


# ═════════════════════════════════════════════════════════════════════
# SIDEBAR
# ═════════════════════════════════════════════════════════════════════
def _sidebar_ativos(nonce: int) -> None:
    with st.sidebar.expander("⚙️ Configurações & Lista de Ativos", expanded=False):
        st.markdown("### Lista de Ativos (Merge)")
        upload = st.file_uploader(
            "Upload de Contingência de Ativos",
            type=["csv", "xlsx", "xlsm"],
            key=f"uploaded_ativos_manual_{nonce}",
        )
        if upload:
            try:
                b = upload.getvalue()
                hid = hash_upload(upload.name, b)
                if st.session_state.get("arquivo_ativos_processado") != hid:
                    raw = DataLoader.ler_arquivo(b, upload.name)
                    proc = DataLoader.processar_lista_ativos(raw)
                    if proc.empty:
                        raise ValueError(
                            "A base de ativos não possui coluna LOGIN válida."
                        )
                    st.session_state["df_gs_manual"] = proc
                    st.session_state["arquivo_ativos_processado"] = hid
                    st.toast(f"✅ {len(proc)} ativos carregados.", icon="👥")
            except Exception as erro:
                st.error(f"Erro ao processar base de ativos: {erro}")

        if st.button(
            "🔄 Reiniciar Aplicação", use_container_width=True, type="secondary"
        ):
            limpar_sessao(incluir_ativos=True)
            st.cache_data.clear()
            st.rerun()
        if st.button("🗑️ Limpar Cache", use_container_width=True, type="secondary"):
            limpar_sessao(incluir_ativos=True)
            st.cache_data.clear()
            st.rerun()


def _sidebar_filtros() -> tuple[float, float, float, float]:
    st.sidebar.markdown("---")
    st.sidebar.markdown(
        '<div style="font-size:12px;font-weight:700;color:#64748B;text-transform:uppercase;">🎯 Filtros de Projeção</div>',
        unsafe_allow_html=True,
    )
    p_ot = st.sidebar.slider("Prob. Otimista de Quebra (%)", 0, 100, 15, step=5) / 100
    p_base = st.sidebar.slider("Prob. Base de Quebra (%)", 0, 100, 30, step=5) / 100
    p_pess = (
        st.sidebar.slider("Prob. Pessimista de Quebra (%)", 0, 100, 50, step=5) / 100
    )
    min_aloc = float(
        st.sidebar.number_input("Mínimo de Alocações", min_value=1, value=5)
    )

    if not (p_ot <= p_base <= p_pess):
        st.sidebar.warning("Atenção: Otimista ≤ Base ≤ Pessimista é o esperado.")

    return p_ot, p_base, p_pess, min_aloc


def _sidebar_robo(df_ativos: pd.DataFrame) -> None:
    st.sidebar.markdown("---")
    if not (ROBO_DISPONIVEL and _ROBO_FN):
        st.sidebar.info("🤖 Robô offline. Utilize upload manual.")
        return

    try:
        _ROBO_FN(
            etl_fn=lambda df_raw, df_gs_arg=None, **_: DataLoader.callback_robo_etl(
                df_raw,
                df_gs_arg if df_gs_arg is not None else df_ativos,
            ),
            gsheets_fn=lambda: df_ativos,
            pasta_padrao=st.session_state.get(
                "robo_pasta_alvo", str(Path.home() / "Downloads")
            ),
            ciclos_estabilidade=1,
            mostrar_toggle=True,
            mostrar_config=True,
        )
    except Exception as erro:
        st.sidebar.error(f"Erro ao inicializar Robô: {erro}")


# ═════════════════════════════════════════════════════════════════════
# UPLOAD PRINCIPAL
# ═════════════════════════════════════════════════════════════════════
def _processar_upload(upload: Any, df_ativos: pd.DataFrame) -> None:
    if not upload:
        return
    try:
        b = upload.getvalue()
        hid = hash_upload(upload.name, b)
        if st.session_state.get("arquivo_processado") == hid:
            return

        with st.spinner(f"Processando {upload.name}..."):
            raw = DataLoader.ler_arquivo(b, upload.name)
            base = DataLoader.preparar_base(raw, df_ativos, filename=upload.name)
            if base.empty:
                raise ValueError("Nenhuma O.S. válida permaneceu após o tratamento.")

            st.session_state.update(
                {
                    "df_memoria_raw": raw,
                    "df_memoria": base,
                    "arquivo_processado": hid,
                    "origem_dados": f"Upload ({upload.name})",
                    "_robo_dados_disponiveis": True,
                    "mensagem_flash": "✅ Base processada com sucesso.",
                }
            )
            st.rerun()
    except Exception as erro:
        st.error(f"❌ Erro no processamento do arquivo: {erro}")


def _render_upload_os(dados_prontos: bool, nonce: int) -> Any:
    if not dados_prontos:
        return st.file_uploader(
            "📂 Carregar base de O.S. manualmente",
            type=["csv", "xlsx", "xlsm"],
            key=f"upload_os_principal_{nonce}",
        )
    with st.expander("📂 Substituir base de O.S. manualmente", expanded=False):
        return st.file_uploader(
            "Upload Manual O.S.",
            type=["csv", "xlsx", "xlsm"],
            key=f"upload_os_substituir_{nonce}",
        )


# ═════════════════════════════════════════════════════════════════════
# MAIN
# ═════════════════════════════════════════════════════════════════════
def main() -> None:
    injetar_css()

    msg = st.session_state.pop("mensagem_flash", "")
    if msg:
        st.success(str(msg))

    render_sidebar_brand(
        titulo="TOTALE",
        subtitulo="Quebra Operacional",
        logo="monitoring",
        ambiente="produção",
        versao="v5.4.0",
        mostrar_data=True,
    )

    if "upload_nonce" not in st.session_state:
        st.session_state["upload_nonce"] = 0
    nonce = int(st.session_state["upload_nonce"])

    _sidebar_ativos(nonce)

    # Base de Ativos
    df_ativos_sessao = df_sessao("df_gs_manual")
    if df_ativos_sessao is not None and not df_ativos_sessao.empty:
        df_ativos = df_ativos_sessao
    else:
        df_ativos = DataLoader.buscar_gsheets()
        if df_ativos is None:
            df_ativos = pd.DataFrame()

    if not df_ativos.empty:
        st.sidebar.caption(f"👥 Base de Ativos: {len(df_ativos)} registros carregados")
    else:
        st.sidebar.warning(
            "⚠️ Lista de Ativos indisponível. Utilize upload manual em Configurações."
        )

    _sidebar_robo(df_ativos)

    df_atual = df_sessao("df_memoria")
    dados_prontos = df_atual is not None and not df_atual.empty

    if render_bloco_importacao(dados_prontos=dados_prontos):
        st.session_state["mensagem_flash"] = "✅ Relatório atualizado com sucesso."
        st.rerun()

    upload = _render_upload_os(dados_prontos, nonce)
    _processar_upload(upload, df_ativos)

    df = df_sessao("df_memoria")
    if df is None or df.empty:
        st.info("Carregue uma base de O.S. para iniciar a análise.")
        return

    p_ot, p_base, p_pess, min_aloc = _sidebar_filtros()

    regioes = (
        sorted(df["REGIÃO"].dropna().astype(str).unique().tolist())
        if "REGIÃO" in df.columns
        else ["TODAS"]
    )
    render_hero(
        "Super Relatório Corporativo",
        "Análise unificada de desempenho operacional",
        regioes,
        len(df),
        badge="TOTALE OPERACIONAL",
        origem=str(st.session_state.get("origem_dados", "Base Carregada")),
    )

    if "TÉCNICO" in df.columns:
        nao_mapeados = int(df["TÉCNICO"].eq("NÃO MAPEADO").sum())
        if nao_mapeados > 0:
            pct = nao_mapeados / len(df) * 100
            st.warning(
                f"⚠️ **Aviso de Integração:** {nao_mapeados} O.S. ({pct:.1f}% do volume ativo) "
                "não possuem correspondência válida na Lista de Ativos."
            )
        else:
            st.success("🎉 Todos os logins estão mapeados na Lista de Ativos.")

    aba = st.radio(
        "Navegação Principal",
        ["Resumo Executivo", "Análise Detalhada", "Auditoria"],
        horizontal=True,
    )

    rotas: dict[str, Callable[[], None]] = {
        "Resumo Executivo": lambda: view_resumo_executivo(df, Config.SLA_QUEBRA_MAXIMA),
        "Análise Detalhada": lambda: view_analise_detalhada(
            df, p_ot, p_base, p_pess, min_aloc, Config.SLA_QUEBRA_MAXIMA
        ),
        "Auditoria": lambda: view_auditoria(df),
    }
    rotas[aba]()

    st.divider()


if __name__ == "__main__":
    main()
